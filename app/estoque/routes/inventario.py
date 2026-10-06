"""Conferência física de munições por embalagem."""
from datetime import datetime
from io import BytesIO

from flask import flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app import db
from app.estoque import estoque_bp
from app.estoque.inventario_models import ConferenciaInventario, ConferenciaInventarioItem
from app.estoque.models import ItemEstoque
from app.models import Configuracao
from app.produtos.models import Produto
from app.produtos.configs.models import TipoProduto
from app.utils.datetime import now_local


def _config(chave, padrao=""):
    registro = Configuracao.query.filter_by(chave=chave).first()
    return (registro.valor if registro and registro.valor is not None else padrao).strip()


def _dados_loja():
    return {
        "loja_cnpj": _config("loja_cnpj"),
        "loja_nome": _config("loja_nome_fantasia", _config("loja_razao_social", "")),
        "loja_endereco": _config("loja_endereco"),
        "loja_cr": _config("loja_cr"),
    }


def _produtos_municao():
    query = Produto.query.outerjoin(TipoProduto, Produto.tipo_id == TipoProduto.id)
    produtos = query.filter(TipoProduto.nome.ilike("%muni%" )).order_by(Produto.nome).all()
    return produtos or Produto.query.order_by(Produto.nome).all()


def _produto_payload(produto, lote="", identificacao="", quantidade=0):
    return {
        "produto_id": produto.id,
        "codigo": produto.codigo or "",
        "descricao": produto.nome,
        "calibre": produto.calibre_rel.nome if produto.calibre_rel else "",
        "lote": lote or "",
        "identificacao_embalagem": identificacao or "",
        "quantidade_embalagem": quantidade or 0,
    }


@estoque_bp.route("/inventario-municoes", methods=["GET", "POST"])
@login_required
def inventario_municoes():
    if request.method == "POST":
        data_raw = request.form.get("data_conferencia") or now_local().date().isoformat()
        try:
            data_conferencia = datetime.strptime(data_raw, "%Y-%m-%d").date()
        except ValueError:
            flash("Data de conferência inválida.", "danger")
            return redirect(url_for("estoque.inventario_municoes"))
        conferencia = ConferenciaInventario(
            data_conferencia=data_conferencia,
            observacoes=request.form.get("observacoes", "").strip() or None,
            **_dados_loja(),
        )
        db.session.add(conferencia)
        db.session.commit()
        return redirect(url_for("estoque.inventario_municoes_detalhe", conferencia_id=conferencia.id))
    conferencias = ConferenciaInventario.query.order_by(ConferenciaInventario.data_conferencia.desc(), ConferenciaInventario.id.desc()).all()
    return render_template("estoque/municoes/inventario.html", conferencias=conferencias, hoje=now_local().date())


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>")
@login_required
def inventario_municoes_detalhe(conferencia_id):
    conferencia = ConferenciaInventario.query.get_or_404(conferencia_id)
    produtos = _produtos_municao()
    return render_template("estoque/municoes/inventario_detalhe.html", conferencia=conferencia, produtos=produtos)


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>/scan", methods=["POST"])
@login_required
def inventario_municoes_scan(conferencia_id):
    conferencia = ConferenciaInventario.query.get_or_404(conferencia_id)
    if conferencia.status != "aberta":
        return jsonify(success=False, message="Esta conferência já foi finalizada."), 409
    payload = request.get_json(silent=True) or {}
    codigo_bipado = str(payload.get("codigo") or payload.get("identificacao_embalagem") or "").strip()
    if not codigo_bipado:
        return jsonify(success=False, message="Informe ou leia a identificação da embalagem."), 400
    if ConferenciaInventarioItem.query.filter_by(conferencia_id=conferencia.id, identificacao_embalagem=codigo_bipado).first():
        return jsonify(success=False, message="Esta embalagem já foi lida nesta conferência."), 409

    estoque = ItemEstoque.query.filter_by(numero_embalagem=codigo_bipado, tipo_item="municao").filter(ItemEstoque.status == "disponivel").first()
    produto_id = payload.get("produto_id") or (estoque.produto_id if estoque else None)
    produto = Produto.query.get(produto_id) if produto_id else None
    if not produto:
        return jsonify(success=False, message="Selecione a munição para esta embalagem."), 400
    lote = str(payload.get("lote") or (estoque.lote if estoque else "")).strip().upper()
    try:
        quantidade = int(payload.get("quantidade_embalagem") or (estoque.quantidade if estoque else 0))
    except (TypeError, ValueError):
        quantidade = 0
    if quantidade <= 0:
        return jsonify(success=False, message="Informe a quantidade de munições da embalagem."), 400

    item = ConferenciaInventarioItem(
        conferencia_id=conferencia.id,
        produto_id=produto.id,
        codigo_municao=produto.codigo or "",
        descricao=produto.nome,
        calibre=produto.calibre_rel.nome if produto.calibre_rel else "",
        lote=lote or None,
        identificacao_embalagem=codigo_bipado,
        quantidade_embalagem=quantidade,
        quantidade_total=quantidade,
    )
    db.session.add(item)
    db.session.commit()
    return jsonify(success=True, item={
        "id": item.id,
        "codigo": item.codigo_municao,
        "descricao": item.descricao,
        "calibre": item.calibre or "",
        "lote": item.lote or "",
        "identificacao_embalagem": item.identificacao_embalagem,
        "quantidade_embalagem": item.quantidade_embalagem,
        "quantidade_total": item.quantidade_total,
    })


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>/item/<int:item_id>/excluir", methods=["POST"])
@login_required
def inventario_municoes_item_excluir(conferencia_id, item_id):
    item = ConferenciaInventarioItem.query.filter_by(id=item_id, conferencia_id=conferencia_id).first_or_404()
    conferencia = item.conferencia
    if conferencia.status != "aberta":
        return jsonify(success=False, message="A conferência já foi finalizada."), 409
    db.session.delete(item)
    db.session.commit()
    return jsonify(success=True)


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>/finalizar", methods=["POST"])
@login_required
def inventario_municoes_finalizar(conferencia_id):
    conferencia = ConferenciaInventario.query.get_or_404(conferencia_id)
    if not conferencia.itens:
        flash("Leia ao menos uma embalagem antes de finalizar.", "warning")
    else:
        conferencia.status = "finalizada"
        conferencia.finalizado_em = now_local()
        db.session.commit()
        flash("Conferência finalizada. Os arquivos podem ser exportados.", "success")
    return redirect(url_for("estoque.inventario_municoes_detalhe", conferencia_id=conferencia.id))


def _linhas_exportacao(conferencia):
    agrupados = {}
    for item in conferencia.itens:
        chave = (item.codigo_municao, item.descricao, item.calibre or "", item.lote or "")
        grupo = agrupados.setdefault(chave, {
            "codigo": item.codigo_municao,
            "descricao": item.descricao,
            "calibre": item.calibre or "",
            "lote": item.lote or "",
            "identificacoes": [],
            "embalagens": 0,
            "total": 0,
        })
        grupo["identificacoes"].append(item.identificacao_embalagem)
        grupo["embalagens"] += 1
        grupo["total"] += item.quantidade_total
    return [[
        grupo["codigo"], grupo["descricao"], grupo["calibre"], grupo["lote"],
        ", ".join(grupo["identificacoes"]), grupo["embalagens"], grupo["total"],
    ] for grupo in agrupados.values()]


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>/excel")
@login_required
def inventario_municoes_excel(conferencia_id):
    conferencia = ConferenciaInventario.query.get_or_404(conferencia_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventário Munições"
    ws.append(["CNPJ da Loja", "Nome da Loja", "Endereço da Loja", "CR da Loja", "Data da Conferência"])
    ws.append([conferencia.loja_cnpj, conferencia.loja_nome, conferencia.loja_endereco, conferencia.loja_cr, conferencia.data_conferencia.strftime("%d/%m/%Y")])
    ws.append([])
    headers = ["Código da Munição", "Descrição", "Calibre", "Lote", "Identificações das Embalagens", "Quantidade de Embalagens", "Quantidade Total da Munição"]
    ws.append(headers)
    for row in _linhas_exportacao(conferencia):
        ws.append(row)
    for cell in ws[1] + ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 14), 42)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"inventario_municoes_{conferencia.data_conferencia.isoformat()}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@estoque_bp.route("/inventario-municoes/<int:conferencia_id>/pdf")
@login_required
def inventario_municoes_pdf(conferencia_id):
    conferencia = ConferenciaInventario.query.get_or_404(conferencia_id)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("INVENTÁRIO FÍSICO DE MUNIÇÕES", styles["Title"]), Spacer(1, 4 * mm)]
    story.append(Paragraph(f"<b>Loja:</b> {conferencia.loja_nome or '-'} &nbsp;&nbsp; <b>CNPJ:</b> {conferencia.loja_cnpj or '-'} &nbsp;&nbsp; <b>CR:</b> {conferencia.loja_cr or '-'}", styles["Normal"]))
    story.append(Paragraph(f"<b>Endereço:</b> {conferencia.loja_endereco or '-'} &nbsp;&nbsp; <b>Data da conferência:</b> {conferencia.data_conferencia.strftime('%d/%m/%Y')}", styles["Normal"]))
    story.append(Spacer(1, 5 * mm))
    rows = [["Código", "Descrição", "Calibre", "Lote", "Identificações das embalagens", "Qtd. embalagens", "Qtd. total"]]
    rows.extend(_linhas_exportacao(conferencia))
    table = Table(rows, repeatRows=1, colWidths=[27 * mm, 75 * mm, 25 * mm, 25 * mm, 55 * mm, 25 * mm, 22 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph(f"Total de embalagens conferidas: <b>{len(conferencia.itens)}</b> &nbsp;&nbsp; Total de munições: <b>{sum(i.quantidade_total for i in conferencia.itens)}</b>", styles["Normal"]))
    doc.build(story)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"inventario_municoes_{conferencia.data_conferencia.isoformat()}.pdf", mimetype="application/pdf")
